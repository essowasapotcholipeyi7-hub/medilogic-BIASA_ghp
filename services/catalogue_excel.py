# -*- coding: utf-8 -*-
"""Catalogue (actes, produits) : export Excel et import Excel avec aperçu — étape 3 de l'abandon de
Google Sheets (patron, 2026-10-18). Le catalogue vit dans Neon (services/feuilles_neon.py) ; on garde le
travail « dans un tableur » par fichier Excel.

Import : rapproché par ID ; seules les cellules qui changent sont écrites ; les lignes sans ID sont
ajoutées (ID suivant) ; une ligne du logiciel absente du fichier est laissée telle quelle (jamais de
suppression par import) ; colonnes inconnues ignorées ; ID, structure_id et (produits) quantite_stock ne
sont jamais modifiés par un import (le stock bouge par les entrées et l'inventaire, tracés)."""
from datetime import date, datetime
from io import BytesIO

from services.feuilles_neon import FeuilleNeon, _texte_cellule

COLONNES_PROTEGEES = {'actes': {'ID', 'structure_id'}, 'produits': {'ID', 'structure_id', 'quantite_stock'}}
COLONNES_SENSIBLES = {'pbr', 'PBR AMU', 'prise_en_charge_amu', 'prise_en_charge_cac', 'AMU-TNS', 'statut'}


def titre_feuille(structure_id, genre):
    return f"struct_{structure_id}_{genre}"


def _nombre(t):
    try:
        return float(str(t).strip().replace(' ', '').replace(' ', '').replace(',', '.'))
    except ValueError:
        return None


def _identiques(a, b):
    a, b = str(a or '').strip(), str(b or '').strip()
    if a == b:
        return True
    na, nb = _nombre(a), _nombre(b)
    return na is not None and nb is not None and abs(na - nb) < 1e-9


def _texte_excel(v):
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d') if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime('%Y-%m-%d %H:%M')
    if isinstance(v, date):
        return v.strftime('%Y-%m-%d')
    return _texte_cellule(v).strip()


# ------------------------------------------------------------------ export
def exporter(structure_id, genre):
    """Fichier .xlsx du catalogue (toutes les colonnes de la structure)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    valeurs = FeuilleNeon(titre_feuille(structure_id, genre)).get_all_values()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Actes' if genre == 'actes' else 'Produits'
    for i, ligne in enumerate(valeurs):
        if i == 0:
            ws.append(ligne)
            continue
        ws.append([(n if (n := _nombre(c)) is not None and str(c).strip() and not str(c).strip().startswith('0')
                    or str(c).strip() == '0' else c) if c != '' else None for c in ligne])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0A3D5C')
    ws.freeze_panes = 'C2'
    for j, entete in enumerate(valeurs[0] if valeurs else [], start=1):
        largeur = max([len(str(entete))] + [len(str(r[j - 1])) for r in valeurs[1:200] if len(r) >= j])
        ws.column_dimensions[get_column_letter(j)].width = min(max(largeur + 2, 8), 60)
    sortie = BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


# ------------------------------------------------------------------ import
def analyser(structure_id, genre, contenu):
    """Compare le fichier Excel au catalogue. Renvoie un dict (rien n'est écrit)."""
    from openpyxl import load_workbook
    feuille = FeuilleNeon(titre_feuille(structure_id, genre))
    grille = feuille.get_all_values()
    if not grille:
        raise ValueError("Le catalogue de la structure est vide ou introuvable.")
    entetes = [e.strip() for e in grille[0]]
    try:
        wb = load_workbook(BytesIO(contenu), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Fichier illisible : enregistrez-le au format Excel (.xlsx).")
    ws = wb.worksheets[0]
    lignes = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    while lignes and not any(c not in (None, '') for c in lignes[0]):
        lignes.pop(0)
    if not lignes:
        raise ValueError("Le fichier est vide.")
    entetes_fichier = [_texte_excel(c) for c in lignes[0]]
    par_nom = {e.lower(): i for i, e in enumerate(entetes) if e}
    correspondance, ignorees = {}, []
    for j, e in enumerate(entetes_fichier):
        if not e:
            continue
        i = par_nom.get(e.lower())
        if i is None:
            ignorees.append(e)
        else:
            correspondance[j] = i
    if 'id' not in {entetes_fichier[j].lower() for j in correspondance} or 'nom' not in {entetes_fichier[j].lower() for j in correspondance}:
        raise ValueError("Le fichier doit avoir au moins les colonnes « ID » et « nom » (partez d'un export du catalogue).")
    j_id = next(j for j in correspondance if entetes_fichier[j].lower() == 'id')
    protegees = COLONNES_PROTEGEES[genre]

    rang_par_id = {}
    for rang, ligne in enumerate(grille[1:], start=2):
        ident = str(ligne[0]).strip() if ligne else ''
        if ident and ident not in rang_par_id:
            rang_par_id[ident] = rang
    modifications, nouveaux, vus, erreurs = [], [], set(), []
    for n, ligne in enumerate(lignes[1:], start=2):
        textes = {correspondance[j]: _texte_excel(ligne[j] if j < len(ligne) else None) for j in correspondance}
        if not any(textes.values()):
            continue
        ident = textes.get(0 if entetes[0] == 'ID' else entetes.index('ID'), '')
        ident = str(int(_nombre(ident))) if _nombre(ident) is not None and float(_nombre(ident)).is_integer() else ident
        nom = textes.get(entetes.index('nom'), '')
        if ident and ident in vus:
            erreurs.append(f"Ligne {n} du fichier : ID {ident} présent deux fois — seule la 1re est prise.")
            continue
        if ident:
            vus.add(ident)
        if ident and ident in rang_par_id:
            rang = rang_par_id[ident]
            existante = grille[rang - 1]
            for i, nouveau in textes.items():
                col = entetes[i]
                if col in protegees:
                    continue
                ancien = existante[i] if i < len(existante) else ''
                if not _identiques(ancien, nouveau):
                    modifications.append({'id': ident, 'nom': existante[1] if len(existante) > 1 else nom,
                                          'colonne': col, 'avant': ancien, 'apres': nouveau, 'rang': rang, 'index': i,
                                          'sensible': col in COLONNES_SENSIBLES})
        else:
            if not nom:
                erreurs.append(f"Ligne {n} du fichier : pas de nom — ignorée.")
                continue
            nouveaux.append({'id': ident, 'nom': nom, 'valeurs': {entetes[i]: v for i, v in textes.items()
                                                                  if entetes[i] not in ('ID', 'structure_id')}})
    absents = len([i for i in rang_par_id if i not in vus])
    return {'entetes': entetes, 'modifications': modifications, 'nouveaux': nouveaux, 'ignorees': ignorees,
            'erreurs': erreurs, 'absents': absents, 'total_catalogue': len(rang_par_id)}


def appliquer(structure_id, genre, contenu):
    """Recalcule l'aperçu (jamais confiance à ce que renvoie le navigateur) puis écrit."""
    res = analyser(structure_id, genre, contenu)
    feuille = FeuilleNeon(titre_feuille(structure_id, genre))
    entetes = res['entetes']
    ecritures = {}
    for m in res['modifications']:
        ecritures[(m['rang'], m['index'] + 1)] = m['apres']
    if ecritures:
        feuille._ecrire_cellules(ecritures)
    ids = [int(_nombre(l[0])) for l in feuille.get_all_values()[1:] if l and _nombre(l[0]) is not None]
    prochain = (max(ids) if ids else 0) + 1
    ajoutes = []
    for nv in res['nouveaux']:
        ident = nv['id'] if nv['id'] and str(nv['id']).isdigit() else str(prochain)
        if not (nv['id'] and str(nv['id']).isdigit()):
            prochain += 1
        ligne = []
        for e in entetes:
            if e == 'ID':
                ligne.append(ident)
            elif e == 'structure_id':
                ligne.append(str(structure_id))
            else:
                ligne.append(nv['valeurs'].get(e, ''))
        feuille.append_row(ligne)
        ajoutes.append({'id': ident, 'nom': nv['nom']})
    return {'modifications': len(res['modifications']), 'ajoutes': ajoutes, 'erreurs': res['erreurs']}
